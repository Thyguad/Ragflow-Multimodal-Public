#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

import json
import logging
import os
import re
import shutil
from datetime import datetime
import html as html_lib
import unicodedata
from functools import reduce
from io import BytesIO
from pathlib import Path
from timeit import default_timer as timer

from bs4 import BeautifulSoup
from docx import Document
from docx.image.exceptions import InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError
from markdown import markdown 
from PIL import Image
from tika import parser

from api.db import LLMType
from api.db.services.llm_service import LLMBundle
from deepdoc.parser import DocxParser, ExcelParser, HtmlParser, JsonParser, MarkdownParser, PdfParser, TxtParser
from deepdoc.parser.figure_parser import VisionFigureParser, vision_figure_parser_figure_data_wraper
from deepdoc.parser.pdf_parser import PlainParser, VisionParser
from rag.image_asset_utils import (
    captions_hash,
    compute_sha256_bytes,
    compute_sha256_file,
    derive_run_id,
    ensure_image_output_root,
    get_doc_asset_dir,
    get_doc_captions_path,
    get_doc_figures_dir,
    get_doc_manifest_path,
)
from rag.nlp import concat_img, find_codec, naive_merge, naive_merge_with_images, naive_merge_docx, rag_tokenizer, tokenize_chunks, tokenize_chunks_with_images, tokenize_table
from rag.utils import num_tokens_from_string


class Docx(DocxParser):
    def __init__(self):
        pass

    def get_picture(self, document, paragraph):
        img = paragraph._element.xpath('.//pic:pic')
        if not img:
            return None
        img = img[0]
        embed = img.xpath('.//a:blip/@r:embed')
        if not embed:
            return None
        embed = embed[0]
        related_part = document.part.related_parts[embed]
        try:
            image_blob = related_part.image.blob
        except UnrecognizedImageError:
            logging.info("Unrecognized image format. Skipping image.")
            return None
        except UnexpectedEndOfFileError:
            logging.info("EOF was unexpectedly encountered while reading an image stream. Skipping image.")
            return None
        except InvalidImageStreamError:
            logging.info("The recognized image stream appears to be corrupted. Skipping image.")
            return None
        except UnicodeDecodeError:
            logging.info("The recognized image stream appears to be corrupted. Skipping image.")
            return None
        try:
            image = Image.open(BytesIO(image_blob)).convert('RGB')
            return image
        except Exception:
            return None

    def __clean(self, line):
        line = re.sub(r"\u3000", " ", line).strip()
        return line

    def __get_nearest_title(self, table_index, filename):
        """Get the hierarchical title structure before the table"""
        import re
        from docx.text.paragraph import Paragraph
        
        titles = []
        blocks = []
        
        # Get document name from filename parameter
        doc_name = re.sub(r"\.[a-zA-Z]+$", "", filename)
        if not doc_name:
            doc_name = "Untitled Document"
            
        # Collect all document blocks while maintaining document order
        try:
            # Iterate through all paragraphs and tables in document order
            for i, block in enumerate(self.doc._element.body):
                if block.tag.endswith('p'):  # Paragraph
                    p = Paragraph(block, self.doc)
                    blocks.append(('p', i, p))
                elif block.tag.endswith('tbl'):  # Table
                    blocks.append(('t', i, None))  # Table object will be retrieved later
        except Exception as e:
            logging.error(f"Error collecting blocks: {e}")
            return ""
            
        # Find the target table position
        target_table_pos = -1
        table_count = 0
        for i, (block_type, pos, _) in enumerate(blocks):
            if block_type == 't':
                if table_count == table_index:
                    target_table_pos = pos
                    break
                table_count += 1
                
        if target_table_pos == -1:
            return ""  # Target table not found
            
        # Find the nearest heading paragraph in reverse order
        nearest_title = None
        for i in range(len(blocks)-1, -1, -1):
            block_type, pos, block = blocks[i]
            if pos >= target_table_pos:  # Skip blocks after the table
                continue
                
            if block_type != 'p':
                continue
                
            if block.style and re.search(r"Heading\s*(\d+)", block.style.name, re.I):
                try:
                    level_match = re.search(r"(\d+)", block.style.name)
                    if level_match:
                        level = int(level_match.group(1))
                        if level <= 7:  # Support up to 7 heading levels
                            title_text = block.text.strip()
                            if title_text:  # Avoid empty titles
                                nearest_title = (level, title_text)
                                break
                except Exception as e:
                    logging.error(f"Error parsing heading level: {e}")
        
        if nearest_title:
            # Add current title
            titles.append(nearest_title)
            current_level = nearest_title[0]
            
            # Find all parent headings, allowing cross-level search
            while current_level > 1:
                found = False
                for i in range(len(blocks)-1, -1, -1):
                    block_type, pos, block = blocks[i]
                    if pos >= target_table_pos:  # Skip blocks after the table
                        continue
                        
                    if block_type != 'p':
                        continue
                        
                    if block.style and re.search(r"Heading\s*(\d+)", block.style.name, re.I):
                        try:
                            level_match = re.search(r"(\d+)", block.style.name)
                            if level_match:
                                level = int(level_match.group(1))
                                # Find any heading with a higher level
                                if level < current_level:  
                                    title_text = block.text.strip()
                                    if title_text:  # Avoid empty titles
                                        titles.append((level, title_text))
                                        current_level = level
                                        found = True
                                        break
                        except Exception as e:
                            logging.error(f"Error parsing parent heading: {e}")
                            
                if not found:  # Break if no parent heading is found
                    break
            
            # Sort by level (ascending, from highest to lowest)
            titles.sort(key=lambda x: x[0])
            # Organize titles (from highest to lowest)
            hierarchy = [doc_name] + [t[1] for t in titles]
            return " > ".join(hierarchy)
            
        return ""

    def __call__(self, filename, binary=None, from_page=0, to_page=100000):
        self.doc = Document(
            filename) if not binary else Document(BytesIO(binary))
        pn = 0
        lines = []
        last_image = None
        for p in self.doc.paragraphs:
            if pn > to_page:
                break
            if from_page <= pn < to_page:
                if p.text.strip():
                    if p.style and p.style.name == 'Caption':
                        former_image = None
                        if lines and lines[-1][1] and lines[-1][2] != 'Caption':
                            former_image = lines[-1][1].pop()
                        elif last_image:
                            former_image = last_image
                            last_image = None
                        lines.append((self.__clean(p.text), [former_image], p.style.name))
                    else:
                        current_image = self.get_picture(self.doc, p)
                        image_list = [current_image]
                        if last_image:
                            image_list.insert(0, last_image)
                            last_image = None
                        lines.append((self.__clean(p.text), image_list, p.style.name if p.style else ""))
                else:
                    if current_image := self.get_picture(self.doc, p):
                        if lines:
                            lines[-1][1].append(current_image)
                        else:
                            last_image = current_image
            for run in p.runs:
                if 'lastRenderedPageBreak' in run._element.xml:
                    pn += 1
                    continue
                if 'w:br' in run._element.xml and 'type="page"' in run._element.xml:
                    pn += 1
        new_line = [(line[0], reduce(concat_img, line[1]) if line[1] else None) for line in lines]

        tbls = []
        for i, tb in enumerate(self.doc.tables):
            title = self.__get_nearest_title(i, filename)
            html = "<table>"
            if title:
                html += f"<caption>Table Location: {title}</caption>"
            for r in tb.rows:
                html += "<tr>"
                i = 0
                while i < len(r.cells):
                    span = 1
                    c = r.cells[i]
                    for j in range(i + 1, len(r.cells)):
                        if c.text == r.cells[j].text:
                            span += 1
                            i = j
                        else:
                            break
                    i += 1
                    html += f"<td>{c.text}</td>" if span == 1 else f"<td colspan='{span}'>{c.text}</td>"
                html += "</tr>"
            html += "</table>"
            tbls.append(((None, html), ""))
        return new_line, tbls


class Pdf(PdfParser):
    def __init__(self):
        super().__init__()

    def __call__(self, filename, binary=None, from_page=0,
                 to_page=100000, zoomin=3, callback=None, separate_tables_figures=False):
        start = timer()
        first_start = start
        callback(msg="OCR started")
        self.__images__(
            filename if not binary else binary,
            zoomin,
            from_page,
            to_page,
            callback
        )
        callback(msg="OCR finished ({:.2f}s)".format(timer() - start))
        logging.info("OCR({}~{}): {:.2f}s".format(from_page, to_page, timer() - start))

        start = timer()
        self._layouts_rec(zoomin)
        callback(0.63, "Layout analysis ({:.2f}s)".format(timer() - start))

        start = timer()
        self._table_transformer_job(zoomin)
        callback(0.65, "Table analysis ({:.2f}s)".format(timer() - start))

        start = timer()
        self._text_merge()
        callback(0.67, "Text merged ({:.2f}s)".format(timer() - start))

        if separate_tables_figures:
            tbls, figures = self._extract_table_figure(True, zoomin, True, True, True)
            self._concat_downward()
            logging.info("layouts cost: {}s".format(timer() - first_start))
            return [(b["text"], self._line_tag(b, zoomin)) for b in self.boxes], tbls, figures
        else:
            tbls = self._extract_table_figure(True, zoomin, True, True)
            # self._naive_vertical_merge()
            self._concat_downward()
            # self._filter_forpages()
            logging.info("layouts cost: {}s".format(timer() - first_start))
            return [(b["text"], self._line_tag(b, zoomin)) for b in self.boxes], tbls


class Markdown(MarkdownParser):
    def get_picture_urls(self, sections):
        if not sections:
            return []
        if isinstance(sections, type("")):
            text = sections
        elif isinstance(sections[0], type("")):
            text = sections[0]
        else:
            return []
        
        from bs4 import BeautifulSoup
        html_content = markdown(text)
        soup = BeautifulSoup(html_content, 'html.parser')
        html_images = [img.get('src') for img in soup.find_all('img') if img.get('src')]
        return html_images
    
    def get_pictures(self, text):
        """Download and open all images from markdown text."""
        import requests
        image_urls = self.get_picture_urls(text)
        images = []
        # Find all image URLs in text
        for url in image_urls:
            try:
                response = requests.get(url, stream=True, timeout=30)
                if response.status_code == 200 and response.headers['Content-Type'].startswith('image/'):
                    img = Image.open(BytesIO(response.content)).convert('RGB')
                    images.append(img)
            except Exception as e:
                logging.error(f"Failed to download/open image from {url}: {e}")
                continue
                    
        return images if images else None

    def __call__(self, filename, binary=None):
        if binary:
            encoding = find_codec(binary)
            txt = binary.decode(encoding, errors="ignore")
        else:
            with open(filename, "r") as f:
                txt = f.read()
        remainder, tables = self.extract_tables_and_remainder(f'{txt}\n')
        sections = []
        tbls = []
        for sec in remainder.split("\n"):
            if num_tokens_from_string(sec) > 3 * self.chunk_token_num:
                sections.append((sec[:int(len(sec) / 2)], ""))
                sections.append((sec[int(len(sec) / 2):], ""))
            else:
                if sec.strip().find("#") == 0:
                    sections.append((sec, ""))
                elif sections and sections[-1][0].strip().find("#") == 0:
                    sec_, _ = sections.pop(-1)
                    sections.append((sec_ + "\n" + sec, ""))
                else:
                    sections.append((sec, ""))
        for table in tables:
            tbls.append(((None, markdown(table, extensions=['markdown.extensions.tables'])), ""))
        return sections, tbls


_VLM_IMAGE_REF_PATTERN = re.compile(r'!\[(.*?)\]\((p\d+_\d+)\)')
_VLM_HEADING_PATTERN = re.compile(r'^#{1,6}\s+')
_VLM_LIST_PATTERN = re.compile(r'^(\s*[-*+]\s+|\s*\d+\.\s+)')
_VLM_FIGURE_CAPTION_PATTERN = re.compile(
    r'^(figure|fig\.?|table|tab\.?|图|表)\s*[\dA-Za-z一二三四五六七八九十]',
    re.IGNORECASE,
)
_VLM_TABLE_SEPARATOR_PATTERN = re.compile(
    r'^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)*\|?\s*$'
)
_VLM_FORMULA_EXPLANATION_PATTERN = re.compile(
    r'^\s*(其中|式中|这里|其中，|式中，|where\b|where,|in which\b|denotes\b|表示|可知|即|也就是)',
    re.IGNORECASE,
)
_VLM_FORMULA_INTRO_PATTERN = re.compile(
    r'(公式|表达式|计算式|定义|如下|表示为|可表示成|可表示为|equation|formula|defined as)\s*[:：]?$',
    re.IGNORECASE,
)
_VLM_ABSTRACT_PATTERN = re.compile(r'^\s*(摘要|摘\s*要|abstract)\s*[:：]?', re.IGNORECASE)
_VLM_ABSTRACT_STOP_PATTERN = re.compile(
    r'^\s*(关键词|关键字|key\s*words?|keywords?|中图分类号|文献标识码|收稿日期|基金项目)\s*[:：]?'
    r'|^\s*(?:[1-9]\s*[\.、]\s*)?(引言|绪论|introduction)\b',
    re.IGNORECASE,
)
_VLM_TABLE_NUMERIC_PATTERN = re.compile(
    r'(?<![A-Za-z])[-+]?\d+(?:\.\d+)?\s*(?:%|％|dB|MHz|GHz|kHz|Hz|ms|s|μs|us)?(?![A-Za-z])',
    re.IGNORECASE,
)
_VLM_TABLE_METRIC_PATTERN = re.compile(
    r'(准确率|识别率|正确率|精度|召回|召回率|查准率|查全率|性能|指标|结果|对比|'
    r'accuracy|acc\b|precision|recall|f1|auc|map\b|ap\b|score|rate|result|performance)',
    re.IGNORECASE,
)
_VLM_TABLE_CONTEXT_PATTERN = re.compile(
    r'(相比|优于|高于|提高|提升|性能|准确率|识别率|结果表明|实验结果|鲁棒性|'
    r'I\s*/?\s*Q|IQ|复数|实数|baseline|compared|performance|accuracy|robust)',
    re.IGNORECASE,
)
_VLM_TABLE_VALUE_HEADER_PATTERN = re.compile(r'^(值|取值|数值|value|values?)$', re.IGNORECASE)
_VLM_TABLE_REFERENCE_CELL_PATTERN = re.compile(r'^\s*\[?\d+(?:\s*[-,，]\s*\d+)*\]?\s*$')
_VLM_LATEX_COMMANDS = {
    "alpha": "alpha", "beta": "beta", "gamma": "gamma", "delta": "delta",
    "epsilon": "epsilon", "varepsilon": "epsilon", "zeta": "zeta",
    "eta": "eta", "theta": "theta", "vartheta": "theta", "iota": "iota",
    "kappa": "kappa", "lambda": "lambda", "mu": "mu", "nu": "nu",
    "xi": "xi", "pi": "pi", "rho": "rho", "varrho": "rho",
    "sigma": "sigma", "tau": "tau", "upsilon": "upsilon", "phi": "phi",
    "varphi": "phi", "chi": "chi", "psi": "psi", "omega": "omega",
    "Gamma": "Gamma", "Delta": "Delta", "Theta": "Theta", "Lambda": "Lambda",
    "Xi": "Xi", "Pi": "Pi", "Sigma": "Sigma", "Phi": "Phi", "Psi": "Psi",
    "Omega": "Omega", "sum": "sum", "prod": "prod", "lim": "lim",
    "min": "min", "max": "max", "argmin": "argmin", "argmax": "argmax",
    "log": "log", "ln": "ln", "exp": "exp", "sin": "sin", "cos": "cos",
    "tan": "tan", "sqrt": "sqrt", "times": "*", "cdot": "*",
    "le": "<=", "leq": "<=", "ge": ">=", "geq": ">=", "neq": "!=",
    "approx": "approx", "sim": "approx", "infty": "infinity",
    "rightarrow": "->", "to": "->", "leftarrow": "<-", "pm": "+/-",
}


def _resolve_vlm_config(parser_config):
    # Prefer UI-provided parser_config.vlm_config; fallback to backend env vars.
    vlm_config = parser_config.get("vlm_config") or {}
    if not isinstance(vlm_config, dict):
        vlm_config = {}

    if not vlm_config.get("vlm_api_key"):
        vlm_config["vlm_api_key"] = os.getenv("RAGFLOW_VLM_API_KEY") or os.getenv("DASHSCOPE_API_KEY")
    if not vlm_config.get("vlm_api_url"):
        vlm_config["vlm_api_url"] = os.getenv("RAGFLOW_VLM_API_URL") or "https://dashscope.aliyuncs.com/compatible-mode/v1"
    if not vlm_config.get("vlm_modelid"):
        vlm_config["vlm_modelid"] = os.getenv("RAGFLOW_VLM_MODEL") or "qwen2.5-vl-72b-instruct"

    if not vlm_config.get("markdown_api_key"):
        vlm_config["markdown_api_key"] = os.getenv("RAGFLOW_MD_API_KEY") or vlm_config.get("vlm_api_key")
    if not vlm_config.get("markdown_api_url"):
        vlm_config["markdown_api_url"] = os.getenv("RAGFLOW_MD_API_URL") or vlm_config.get("vlm_api_url")
    if not vlm_config.get("markdown_modelid"):
        vlm_config["markdown_modelid"] = os.getenv("RAGFLOW_MD_MODEL") or "qwen-plus"

    return vlm_config


def _replace_vlm_placeholders(markdown_text, figures_dict, captions_dict, figure_order=None):
    fig_keys = figure_order or list(figures_dict.keys())
    fig_idx = [0]

    def _replace_placeholder(_):
        if fig_idx[0] >= len(fig_keys):
            # Drop unmatched placeholders instead of keeping broken image refs.
            return ""
        key = fig_keys[fig_idx[0]]
        fig_idx[0] += 1
        if not key or key not in figures_dict:
            return ""
        caption = captions_dict.get(key, "图")
        return f"![{caption}]({key})"

    markdown_text = re.sub(r'!\[FIG_PLACEHOLDER\]\(#\)', _replace_placeholder, markdown_text)
    return re.sub(r'\n{3,}', '\n\n', markdown_text).strip()


def _vlm_chunk_token_num(parser_config):
    return max(192, int(parser_config.get("chunk_token_num", 128)))


def _vlm_abstract_token_num(chunk_token_num):
    return min(420, max(int(chunk_token_num), int(chunk_token_num * 2.5)))


def _looks_like_vlm_figure_caption(text):
    return bool(_VLM_FIGURE_CAPTION_PATTERN.match(text.strip()))


def _clean_vlm_heading_text(text):
    text = re.sub(r'^\s*#{1,6}\s*', '', str(text or "")).strip()
    return _normalize_vlm_chunk_text(text) if text else ""


def _vlm_section(text, kind="paragraph", headings=None):
    return (text, {"kind": kind, "headings": list(headings or [])})


def _vlm_section_meta(section):
    if len(section) < 2 or not isinstance(section[1], dict):
        return {"kind": "paragraph", "headings": []}
    return section[1]


def _vlm_heading_path(headings):
    cleaned = []
    for heading in headings or []:
        heading = _normalize_vlm_chunk_text(heading)
        if heading and heading not in cleaned:
            cleaned.append(heading)
    if not cleaned:
        return ""
    return " / ".join(cleaned[-3:])


def _prepend_vlm_heading_context(text, headings):
    text = str(text or "").strip()
    heading_path = _vlm_heading_path(headings)
    if not text or not heading_path:
        return text
    if text.startswith("章节:"):
        return text
    return f"章节: {heading_path}\n{text}"


def _is_vlm_table_separator(line):
    return bool(_VLM_TABLE_SEPARATOR_PATTERN.match(str(line or "")))


def _is_vlm_markdown_table_start(lines, idx):
    if idx + 1 >= len(lines):
        return False
    return "|" in lines[idx] and _is_vlm_table_separator(lines[idx + 1])


def _collect_vlm_markdown_table(lines, idx):
    table_lines = []
    while idx < len(lines):
        stripped = lines[idx].strip()
        if not stripped or "|" not in stripped:
            break
        table_lines.append(stripped)
        idx += 1
    return "\n".join(table_lines), idx


def _collect_vlm_html_table(lines, idx):
    table_lines = []
    while idx < len(lines):
        line = lines[idx].rstrip()
        table_lines.append(line)
        idx += 1
        if "</table>" in line.lower():
            break
    return "\n".join(table_lines), idx


def _looks_like_vlm_formula_explanation(text):
    return bool(_VLM_FORMULA_EXPLANATION_PATTERN.match(str(text or "").strip()))


def _looks_like_vlm_formula_intro(text):
    return bool(_VLM_FORMULA_INTRO_PATTERN.search(str(text or "").strip()))


def _looks_like_vlm_abstract(text, headings=None):
    if _VLM_ABSTRACT_PATTERN.match(str(text or "").strip()):
        return True
    for heading in headings or []:
        if _VLM_ABSTRACT_PATTERN.match(str(heading or "").strip()):
            return True
    return False


def _looks_like_vlm_abstract_stop(text):
    return bool(_VLM_ABSTRACT_STOP_PATTERN.match(str(text or "").strip()))


def _vlm_table_context_snippet(text, max_chars=260):
    text = _normalize_vlm_chunk_text(text)
    text = "\n".join(
        line for line in str(text or "").splitlines()
        if not line.strip().startswith("章节:")
    )
    text = re.sub(r'\s+', ' ', text).strip()
    if len(text) <= max_chars:
        return text
    return text[-max_chars:].strip()


def _looks_like_vlm_table_context(text):
    return bool(_VLM_TABLE_CONTEXT_PATTERN.search(str(text or "")))


def _peek_vlm_table_following_context(lines, idx, max_chars=260):
    collected = []
    chars = 0
    skipped_blank = False
    while idx < len(lines):
        stripped = str(lines[idx] or "").strip()
        if not stripped:
            if collected or skipped_blank:
                break
            skipped_blank = True
            idx += 1
            continue
        if (
            _VLM_HEADING_PATTERN.match(stripped)
            or _is_vlm_markdown_table_start(lines, idx)
            or "<table" in stripped.lower()
            or _VLM_IMAGE_REF_PATTERN.search(stripped)
            or _VLM_LIST_PATTERN.match(stripped)
        ):
            break
        collected.append(stripped)
        chars += len(stripped)
        idx += 1
        if chars >= max_chars:
            break
    return _vlm_table_context_snippet("\n".join(collected), max_chars=max_chars)


def _split_large_vlm_block(text, chunk_token_num):
    text = text.strip()
    if not text:
        return []
    max_section_tokens = max(32, int(chunk_token_num))
    if num_tokens_from_string(text) <= max_section_tokens:
        return [text]

    normalized = re.sub(r'\s*\n\s*', ' ', text)
    sentences = [
        s.strip() for s in re.split(
            r'(?<=[。！？!?；;])\s+|(?<=\.)\s+(?=[A-Z0-9#])',
            normalized,
        ) if s.strip()
    ]
    if len(sentences) <= 1:
        midpoint = max(1, len(normalized) // 2)
        split_at = normalized.rfind(" ", 0, midpoint)
        if split_at <= 0:
            split_at = midpoint
        return [normalized[:split_at].strip(), normalized[split_at:].strip()]

    blocks = []
    current = []
    current_tokens = 0
    for sentence in sentences:
        sentence_tokens = num_tokens_from_string(sentence)
        if current and current_tokens + sentence_tokens > max_section_tokens:
            blocks.append(" ".join(current).strip())
            current = [sentence]
            current_tokens = sentence_tokens
            continue
        current.append(sentence)
        current_tokens += sentence_tokens
    if current:
        blocks.append(" ".join(current).strip())
    return [block for block in blocks if block]


def _build_vlm_sections(markdown_text, chunk_token_num):
    lines = str(markdown_text or "").splitlines()
    sections = []
    tbls = []
    paragraph = []
    heading_stack = []
    recent_text_context = ""
    active_abstract_headings = None
    active_abstract_tokens = 0

    def _current_headings():
        return [heading for _, heading in heading_stack if heading]

    def _flush_paragraph():
        nonlocal paragraph, recent_text_context, active_abstract_headings, active_abstract_tokens
        if not paragraph:
            return
        text = "\n".join(paragraph).strip()
        paragraph = []
        headings = _current_headings()
        abstract_budget = _vlm_abstract_token_num(chunk_token_num)
        text_tokens = num_tokens_from_string(text)
        is_abstract_start = _looks_like_vlm_abstract(text, headings)
        is_abstract_stop = _looks_like_vlm_abstract_stop(text)
        is_abstract_continuation = (
            active_abstract_headings is not None
            and headings == active_abstract_headings
            and not is_abstract_stop
            and active_abstract_tokens + text_tokens <= abstract_budget
        )
        if is_abstract_start or is_abstract_continuation:
            kind = "abstract"
            split_token_num = abstract_budget
            if is_abstract_start:
                active_abstract_headings = list(headings)
                active_abstract_tokens = 0
            active_abstract_tokens += text_tokens
        else:
            active_abstract_headings = None
            active_abstract_tokens = 0
            kind = "formula" if _looks_like_vlm_formula_block(text) else "paragraph"
            split_token_num = chunk_token_num
        for block in _split_large_vlm_block(text, split_token_num):
            sections.append(_vlm_section(block, kind, headings))
        recent_text_context = _vlm_table_context_snippet(text)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            _flush_paragraph()
            i += 1
            continue

        if _VLM_HEADING_PATTERN.match(stripped):
            _flush_paragraph()
            active_abstract_headings = None
            active_abstract_tokens = 0
            heading_level = len(re.match(r'^\s*(#{1,6})', stripped).group(1))
            heading_text = _clean_vlm_heading_text(stripped)
            heading_stack = [
                (level, text) for level, text in heading_stack
                if level < heading_level
            ]
            if heading_text:
                heading_stack.append((heading_level, heading_text))
            i += 1
            continue

        if _is_vlm_markdown_table_start(lines, i):
            _flush_paragraph()
            active_abstract_headings = None
            active_abstract_tokens = 0
            table_text, i = _collect_vlm_markdown_table(lines, i)
            after_context = _peek_vlm_table_following_context(lines, i)
            tbls.extend(
                _build_vlm_table_chunks(
                    [table_text],
                    _current_headings(),
                    contexts=[recent_text_context, after_context],
                )
            )
            continue

        if "<table" in stripped.lower():
            _flush_paragraph()
            active_abstract_headings = None
            active_abstract_tokens = 0
            table_text, i = _collect_vlm_html_table(lines, i)
            after_context = _peek_vlm_table_following_context(lines, i)
            tbls.extend(
                _build_vlm_table_chunks(
                    [table_text],
                    _current_headings(),
                    contexts=[recent_text_context, after_context],
                )
            )
            continue

        if _VLM_IMAGE_REF_PATTERN.search(stripped):
            _flush_paragraph()
            active_abstract_headings = None
            active_abstract_tokens = 0
            image_block = [stripped]
            j = i + 1
            while j < len(lines) and lines[j].strip() and _looks_like_vlm_figure_caption(lines[j].strip()):
                image_block.append(lines[j].strip())
                j += 1
            sections.append(_vlm_section("\n".join(image_block), "image", _current_headings()))
            i = j
            continue

        if _VLM_LIST_PATTERN.match(stripped):
            _flush_paragraph()
            active_abstract_headings = None
            active_abstract_tokens = 0
            list_block = [stripped]
            j = i + 1
            while j < len(lines):
                next_line = lines[j].rstrip()
                next_stripped = next_line.strip()
                if not next_stripped:
                    break
                if _VLM_HEADING_PATTERN.match(next_stripped) or _VLM_IMAGE_REF_PATTERN.search(next_stripped):
                    break
                if _VLM_LIST_PATTERN.match(next_stripped) or next_line.startswith(("  ", "\t")):
                    list_block.append(next_stripped)
                    j += 1
                    continue
                break
            sections.append(_vlm_section("\n".join(list_block), "list", _current_headings()))
            i = j
            continue

        paragraph.append(stripped)
        i += 1

    _flush_paragraph()
    return sections, tbls


def _build_vlm_section_images(sections, figures_dict):
    section_images = []
    for section_text, _ in sections:
        figure_keys = [match[1] for match in _VLM_IMAGE_REF_PATTERN.findall(section_text or "")]
        pil_image = None
        for key in figure_keys:
            if key not in figures_dict:
                continue
            try:
                next_image = Image.open(BytesIO(figures_dict[key])).convert("RGB")
                pil_image = next_image if pil_image is None else concat_img(pil_image, next_image)
            except Exception:
                continue
        section_images.append(pil_image)
    return section_images


_VLM_ZERO_WIDTH_AND_BIDI_RE = re.compile(r"[\u200b\u200c\u200d\u2060\ufeff\u202a-\u202e]")
_VLM_CITATION_ARTIFACT_PATTERNS = (
    re.compile(r"#ID#\s*\d+\s*##\s*\d+\s*(?:\$\s*){1,2}", re.IGNORECASE),
    re.compile(r"#ID#\s*\d+(?:\s*(?:##\s*\d+\s*\$\s*\$|[=#$\s])+)*", re.IGNORECASE),
    re.compile(r"##\s*\d+\s*\$\s*\$"),
    re.compile(r"(?<!\w)##\s*\d+\s*\$\s+\$(?!\w)"),
    re.compile(r"(?<![\w#])\d+\s*==(?=(?:\s|[，。！？；：,.!?;:)\]】》〉”’]|$))"),
)
_VLM_ISOLATED_SYMBOL_RE = re.compile(
    r"(?:(?<=^)|(?<=\s))[①②③④⑤⑥⑦⑧⑨⑩ⓘ○●◆◇■□▲△★☆✦✧❖❀❁❂❃❉❋✪✳✴➤➜➝➞➟➠·•▪▫‣⁃](?=$|\s)"
)


def _clean_vlm_text_noise(text):
    text = str(text or "").replace("ⓘ", " ")
    text = unicodedata.normalize("NFKC", text)
    text = _VLM_ZERO_WIDTH_AND_BIDI_RE.sub("", text)
    text = text.replace("\ufffd", "")
    text = text.replace("\u00a0", " ").replace("\u3000", " ")
    for pattern in _VLM_CITATION_ARTIFACT_PATTERNS:
        text = pattern.sub(" ", text)
    text = _VLM_ISOLATED_SYMBOL_RE.sub(" ", text)
    text = re.sub(r"(?:(?<=^)|(?<=[\s，。！？；：,.!?;:]))i(?=(?:[\s，。！？；：,.!?;:]|$))", " ", text)
    text = re.sub(r"\$\s+\$", "$$", text)
    text = re.sub(r"([(\[{（【《“‘])\s+", r"\1", text)
    text = re.sub(r"\s+([)\]}）》”’】])", r"\1", text)
    text = re.sub(r"\s+([，。！？；：、,.!?;:])", r"\1", text)
    text = re.sub(r"(##\d+\$\$)(?=[A-Za-z\u4e00-\u9fff])", r"\1 ", text)
    return text


def _clean_vlm_latex_markup(text):
    text = str(text or "")
    text = re.sub(r'\\\(([\s\S]*?)\\\)', r'\1', text)
    text = re.sub(r'\\\[([\s\S]*?)\\\]', r'\1', text)
    text = re.sub(r'\$\$([\s\S]*?)\$\$', r'\1', text)
    text = re.sub(r'(?<!\\)\$([^$\n]+?)(?<!\\)\$', r'\1', text)
    for _ in range(3):
        text = re.sub(r'\\frac\s*\{([^{}]+)\}\s*\{([^{}]+)\}', r'(\1)/(\2)', text)
    text = re.sub(r'\\(?:mathrm|mathbf|boldsymbol|text|operatorname)\s*\{([^{}]*)\}', r'\1', text)
    text = re.sub(r'\\(?:left|right)\s*', '', text)
    text = text.replace(r'\%', '%')
    text = text.replace(r'\_', '_')
    text = text.replace(r'\&', '&')

    def _replace_latex_command(match):
        command = match.group(1)
        replacement = _VLM_LATEX_COMMANDS.get(command, command)
        return f" {replacement} "

    text = re.sub(r'\\([A-Za-z]+)\s*', _replace_latex_command, text)
    text = re.sub(
        r'_\s*\{([^{}]+)\}',
        lambda m: "_" + re.sub(r'\s+', "_", m.group(1).strip()),
        text,
    )
    text = re.sub(
        r'\^\s*\{([^{}]+)\}',
        lambda m: "^" + re.sub(r'\s+', "", m.group(1).strip()),
        text,
    )
    text = re.sub(r'\s+([_^])', r'\1', text)
    text = re.sub(r'([_^])\s+([A-Za-z0-9])', r'\1\2', text)
    text = text.replace("{", "").replace("}", "")
    text = text.replace("\\", "")
    text = re.sub(r'[ \t]+', ' ', text)
    return text


def _clean_vlm_inline_html(text):
    text = html_lib.unescape(str(text or ""))
    text = re.sub(
        r'<\s*sub\s*>(.*?)<\s*/\s*sub\s*>',
        lambda m: "_" + re.sub(r'\s+', "_", m.group(1).strip()),
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    text = re.sub(
        r'<\s*sup\s*>(.*?)<\s*/\s*sup\s*>',
        lambda m: "^" + re.sub(r'\s+', "", m.group(1).strip()),
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    text = re.sub(r'<\s*br\s*/?\s*>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<[^>]+>', ' ', text)
    return text


def _looks_like_vlm_formula_block(text):
    text = str(text or "").strip()
    if not text:
        return False
    if re.fullmatch(r'\${1,2}[\s\S]*?\${1,2}', text):
        return True
    if re.search(r'[\u4e00-\u9fff]', text):
        return False
    markers = len(re.findall(r'\\[A-Za-z]+|[_^]\s*\{|[=+\-*/]', text))
    return markers >= 2 and len(text) <= 500


def _normalize_vlm_chunk_text(text):
    text = _clean_vlm_inline_html(text).strip()
    if not text:
        return ""

    text = _VLM_IMAGE_REF_PATTERN.sub(
        lambda m: f"图像说明: {m.group(1).strip()}" if m.group(1).strip() else "图像",
        text,
    )
    text = re.sub(r'^\s*#{1,6}\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    text = re.sub(r'(\*\*|__)(.*?)\1', r'\2', text)
    text = re.sub(r'\*([^*\n]+)\*', r'\1', text)
    text = text.replace("`", "")
    text = _clean_vlm_text_noise(text)
    text = _clean_vlm_latex_markup(text)
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\s*\n\s*', '\n', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def _normalize_vlm_table_text(text):
    text = _clean_vlm_inline_html(text)
    if not text:
        return ""

    text = _clean_vlm_text_noise(text)
    text = _clean_vlm_latex_markup(text)
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    text = re.sub(r'(\*\*|__)(.*?)\1', r'\2', text)
    text = re.sub(r'\*([^*\n]+)\*', r'\1', text)
    text = text.replace("`", "")
    text = re.sub(r'[ \t]+', ' ', text)
    return text.strip()


def _split_vlm_table_cells(line):
    cells = str(line or "").strip()
    if cells.startswith("|"):
        cells = cells[1:]
    if cells.endswith("|"):
        cells = cells[:-1]
    return [_normalize_vlm_table_text(cell.strip()) for cell in cells.split("|")]


def _parse_vlm_markdown_table_rows(table_text):
    rows = []
    for line in str(table_text or "").splitlines():
        stripped = line.strip()
        if not stripped or "|" not in stripped or _is_vlm_table_separator(stripped):
            continue
        cells = [cell for cell in _split_vlm_table_cells(stripped) if cell]
        if cells:
            rows.append(cells)
    return rows


def _parse_vlm_html_table_rows(table_text):
    soup = BeautifulSoup(str(table_text or ""), "html.parser")
    rows = []
    for tr in soup.find_all("tr"):
        cells = []
        for cell in tr.find_all(["th", "td"]):
            value = _normalize_vlm_table_text(cell.get_text(" ", strip=True))
            if value:
                cells.append(value)
        if cells:
            rows.append(cells)
    return rows


def _is_vlm_table_numeric_cell(value):
    value = str(value or "").strip()
    if not value or _VLM_TABLE_REFERENCE_CELL_PATTERN.match(value):
        return False
    return bool(_VLM_TABLE_NUMERIC_PATTERN.search(value))


def _is_vlm_table_metric_label(value):
    return bool(_VLM_TABLE_METRIC_PATTERN.search(str(value or "")))


def _is_vlm_table_value_header(value):
    return bool(_VLM_TABLE_VALUE_HEADER_PATTERN.match(str(value or "").strip()))


def _looks_like_vlm_metric_table(rows):
    if not rows:
        return False
    cells = [cell for row in rows for cell in row if cell]
    if not cells:
        return False
    has_metric_label = any(_is_vlm_table_metric_label(cell) for cell in cells)
    numeric_count = sum(1 for cell in cells if _is_vlm_table_numeric_cell(cell))
    has_percent = any(("%" in cell or "％" in cell) and not _VLM_TABLE_REFERENCE_CELL_PATTERN.match(cell) for cell in cells)
    return has_metric_label and (numeric_count > 0 or has_percent)


def _build_vlm_table_item_lines(rows, max_lines=12):
    if len(rows) < 2:
        return []

    headers = rows[0]
    lines = []
    for row in rows[1:]:
        pairs = []
        for idx, cell in enumerate(row):
            if not cell:
                continue
            header = headers[idx] if idx < len(headers) and headers[idx] else f"列{idx + 1}"
            if header == cell:
                continue
            pairs.append(f"{header}={cell}")
        if pairs:
            lines.append(f"表格项: {'; '.join(pairs)}")
        if len(lines) >= max_lines:
            break
    return lines


def _build_vlm_table_relation_lines(rows, max_lines=12):
    if len(rows) < 2:
        return []

    headers = rows[0]
    lines = []
    seen = set()
    for row in rows[1:]:
        if len(row) < 2:
            continue
        row_label = row[0]
        row_label_is_metric = _is_vlm_table_metric_label(row_label)
        for idx in range(1, min(len(row), len(headers))):
            cell = row[idx]
            col_label = headers[idx]
            if not cell or not col_label or not _is_vlm_table_numeric_cell(cell):
                continue
            relation_label = col_label
            relation_value = cell
            col_label_is_metric = _is_vlm_table_metric_label(col_label)
            if row_label_is_metric:
                relation_label, relation_value = _normalize_vlm_table_relation_label_value(row_label, cell)
                if _is_vlm_table_value_header(col_label):
                    line = f"表格关系: {relation_label} 为 {relation_value}"
                else:
                    line = f"表格关系: {col_label} 的 {relation_label} 为 {relation_value}"
            elif col_label_is_metric:
                relation_label, relation_value = _normalize_vlm_table_relation_label_value(col_label, cell)
                line = f"表格关系: {row_label} 的 {relation_label} 为 {relation_value}"
            else:
                continue
            if line in seen:
                continue
            seen.add(line)
            lines.append(line)
            if len(lines) >= max_lines:
                return lines
    return lines


def _build_vlm_table_context_lines(contexts, max_chars=480):
    if not contexts:
        return []

    selected = []
    seen = set()
    total = 0
    for context in contexts:
        snippet = _vlm_table_context_snippet(context)
        if not snippet or not _looks_like_vlm_table_context(snippet):
            continue
        key = re.sub(r'\s+', '', snippet)
        if key in seen:
            continue
        seen.add(key)
        if total + len(snippet) > max_chars:
            snippet = snippet[:max(0, max_chars - total)].strip()
        if snippet:
            selected.append(snippet)
            total += len(snippet)
        if total >= max_chars:
            break
    if not selected:
        return []
    return [f"相关上下文: {' '.join(selected)}"]


def _normalize_vlm_table_relation_label_value(label, value):
    label = str(label or "").strip()
    value = str(value or "").strip()
    if re.search(r'[/／]\s*[%％]\s*$', label):
        label = re.sub(r'\s*[/／]\s*[%％]\s*$', '', label).strip()
        if value and not value.endswith(("%", "％")):
            value = f"{value}%"
    return label, value


def _vlm_rows_to_table_text(rows, headings=None, contexts=None):
    if not rows:
        return ""

    lines = []
    heading_path = _vlm_heading_path(headings)
    if heading_path:
        lines.append(f"章节: {heading_path}")
    lines.append("表格:")
    for idx, row in enumerate(rows):
        label = "表头" if idx == 0 else "行"
        lines.append(f"{label}: {' | '.join(row)}")
    lines.extend(_build_vlm_table_item_lines(rows))
    if _looks_like_vlm_metric_table(rows):
        lines.extend(_build_vlm_table_relation_lines(rows))
        lines.extend(_build_vlm_table_context_lines(contexts))
        lines.append("关键词: 实验结果, 性能对比, 指标对比, 方法对比, 准确率, 识别率")
    return "\n".join(lines)


def _build_vlm_table_chunks(tables, headings=None, contexts=None):
    tbls = []
    for table in tables:
        raw_table = str(table or "")
        if "<table" in raw_table.lower():
            rows = _parse_vlm_html_table_rows(raw_table)
        else:
            rows = _parse_vlm_markdown_table_rows(raw_table)
        table_text = _vlm_rows_to_table_text(rows, headings, contexts=contexts)
        table_text = _normalize_vlm_table_text(table_text)
        if not table_text:
            continue
        tbls.append(((None, table_text), ""))
    return tbls


def _merge_vlm_sections(sections, chunk_token_num):
    chunks, _ = _merge_vlm_sections_with_images(sections, None, chunk_token_num)
    return chunks


def _vlm_context_tail(text, max_chars=180):
    text = "\n".join(
        line for line in str(text or "").splitlines()
        if not line.strip().startswith("章节:")
    )
    text = re.sub(r'\s+', ' ', text).strip()
    if len(text) <= max_chars:
        return text
    return text[-max_chars:].strip()


def _dedupe_vlm_heading_lines(text):
    lines = []
    seen_headings = set()
    for line in str(text or "").splitlines():
        stripped = line.strip()
        if stripped.startswith("章节:"):
            if stripped in seen_headings:
                continue
            seen_headings.add(stripped)
        lines.append(line)
    return "\n".join(lines).strip()


def _compact_vlm_chunks(chunks, images, chunk_token_num):
    if not chunks:
        return [], []

    min_tokens = max(48, int(chunk_token_num * 0.35))
    max_tokens = max(320, int(chunk_token_num * 2))
    compact_chunks = []
    compact_images = []
    compact_tokens = []

    for chunk, image in zip(chunks, images):
        chunk = str(chunk or "").strip()
        if not chunk:
            continue
        token_count = num_tokens_from_string(chunk)
        if (
            image is None
            and compact_chunks
            and compact_images[-1] is None
            and (compact_tokens[-1] < min_tokens or token_count < min_tokens)
            and compact_tokens[-1] + token_count <= max_tokens
        ):
            compact_chunks[-1] = _dedupe_vlm_heading_lines(f"{compact_chunks[-1]}\n{chunk}")
            compact_tokens[-1] += token_count
            continue
        compact_chunks.append(_dedupe_vlm_heading_lines(chunk))
        compact_images.append(image)
        compact_tokens.append(token_count)

    return compact_chunks, compact_images


def _merge_vlm_sections_with_images(sections, section_images, chunk_token_num):
    if not sections:
        return [], []

    chunks = []
    images = []
    current_parts = []
    current_tokens = 0
    current_formula_only = False
    current_has_formula = False
    current_kind = None
    last_text_context = ""

    def flush():
        nonlocal current_parts, current_tokens, current_formula_only, current_has_formula, current_kind, last_text_context
        if not current_parts:
            return
        chunk_text = "\n".join(part for part in current_parts if part).strip()
        chunk_text = _dedupe_vlm_heading_lines(chunk_text)
        if chunk_text:
            chunks.append(chunk_text)
            images.append(None)
            last_text_context = _vlm_context_tail(chunk_text)
        current_parts = []
        current_tokens = 0
        current_formula_only = False
        current_has_formula = False
        current_kind = None

    if not section_images or len(section_images) != len(sections):
        section_images = [None] * len(sections)

    for (section_text, _), section_image in zip(sections, section_images):
        meta = _vlm_section_meta((section_text, _))
        headings = meta.get("headings", [])
        raw_text = str(section_text or "")
        text = _normalize_vlm_chunk_text(raw_text)
        if not text:
            continue

        section_kind = meta.get("kind", "paragraph")
        text_tokens = num_tokens_from_string(text)
        is_image_caption = bool(_VLM_IMAGE_REF_PATTERN.search(raw_text))
        is_formula_block = section_kind == "formula" or (
            section_kind != "abstract" and _looks_like_vlm_formula_block(raw_text)
        )
        is_formula_explanation = _looks_like_vlm_formula_explanation(raw_text)
        if is_formula_block and not re.match(r'^(公式|Equation)\s*[:：]', text, re.IGNORECASE):
            text = f"公式: {text}"
            text_tokens = num_tokens_from_string(text)
        text = _prepend_vlm_heading_context(text, headings)
        text_tokens = num_tokens_from_string(text)

        if is_image_caption:
            flush()
            if last_text_context and "附近文本:" not in text:
                text = f"{text}\n附近文本: {last_text_context}"
            chunks.append(text)
            images.append(section_image)
            continue

        should_keep_formula_pair = (
            current_has_formula and is_formula_explanation
        ) or (
            is_formula_block and current_parts and _looks_like_vlm_formula_intro(current_parts[-1])
        )
        if current_parts and ((current_kind == "abstract") != (section_kind == "abstract")):
            flush()

        budget = (
            _vlm_abstract_token_num(chunk_token_num)
            if current_parts and current_kind == "abstract" and section_kind == "abstract"
            else chunk_token_num
        )
        if current_parts and current_tokens + text_tokens > budget and not should_keep_formula_pair:
            flush()

        current_parts.append(text)
        current_tokens += text_tokens
        current_kind = section_kind if current_kind is None else current_kind
        current_formula_only = is_formula_block if len(current_parts) == 1 else current_formula_only and is_formula_block
        current_has_formula = current_has_formula or is_formula_block

    flush()
    return _compact_vlm_chunks(chunks, images, chunk_token_num)


def _export_vlm_assets_for_image_recall(filename, figures_dict, captions_dict, binary=None):
    doc_name = os.path.splitext(os.path.basename(filename))[0].strip()
    if not doc_name:
        return

    output_root = ensure_image_output_root()
    doc_output_dir = get_doc_asset_dir(doc_name, output_root)
    figures_dir = get_doc_figures_dir(doc_name, output_root)
    captions_path = get_doc_captions_path(doc_name, output_root)
    manifest_path = get_doc_manifest_path(doc_name, output_root)
    doc_output_dir.mkdir(parents=True, exist_ok=True)

    if figures_dir.exists():
        shutil.rmtree(figures_dir)
    if captions_path.exists():
        captions_path.unlink()

    normalized_captions = {
        key: str(value or "").strip()
        for key, value in (captions_dict or {}).items()
        if key in (figures_dict or {})
    }

    if figures_dict:
        figures_dir.mkdir(parents=True, exist_ok=True)

    for figure_key, figure_bytes in (figures_dict or {}).items():
        if not figure_key or not figure_bytes:
            continue
        figure_path = figures_dir / f"{figure_key}.jpg"
        with open(figure_path, "wb") as fout:
            fout.write(figure_bytes)

    if normalized_captions:
        with open(captions_path, "w", encoding="utf-8") as fout:
            json.dump(normalized_captions, fout, ensure_ascii=False, indent=2)

    source_file_hash = compute_sha256_bytes(binary) if binary else compute_sha256_file(filename)
    manifest = {
        "run_id": derive_run_id(output_root),
        "doc_name": doc_name,
        "source_file_hash": source_file_hash,
        "parser_version": "vlm_image_recall_assets_v1",
        "embedding_model": os.getenv("RAGFLOW_IMAGE_EMBEDDING_MODEL", "qwen3-vl-embedding"),
        "vector_dim": int(os.getenv("RAGFLOW_IMAGE_EMBEDDING_DIMENSION", "1024")),
        "figure_count": len(normalized_captions),
        "caption_hash": captions_hash(normalized_captions),
        "status": "ready" if normalized_captions else "no_figures",
        "created_at": datetime.utcnow().isoformat() + "Z",
        "output_root": str(output_root),
    }
    with open(manifest_path, "w", encoding="utf-8") as fout:
        json.dump(manifest, fout, ensure_ascii=False, indent=2)


def chunk(filename, binary=None, from_page=0, to_page=100000,
          lang="Chinese", callback=None, **kwargs):
    """
        Supported file formats are docx, pdf, excel, txt.
        This method apply the naive ways to chunk files.
        Successive text will be sliced into pieces using 'delimiter'.
        Next, these successive pieces are merge into chunks whose token number is no more than 'Max token number'.
    """

    is_english = lang.lower() == "english"  # is_english(cks)
    parser_config = kwargs.get(
        "parser_config", {
            "chunk_token_num": 128, "delimiter": "\n!?。；！？", "layout_recognize": "DeepDOC"})
    doc = {
        "docnm_kwd": filename,
        "title_tks": rag_tokenizer.tokenize(re.sub(r"\.[a-zA-Z]+$", "", filename))
    }
    doc["title_sm_tks"] = rag_tokenizer.fine_grained_tokenize(doc["title_tks"])
    res = []
    pdf_parser = None
    section_images = None
    if re.search(r"\.docx$", filename, re.IGNORECASE):
        layout_recognizer = parser_config.get("layout_recognize", "DeepDOC")
        if isinstance(layout_recognizer, bool):
            layout_recognizer = "DeepDOC" if layout_recognizer else "Plain Text"

        if layout_recognizer == "VLM":
            from deepdoc.parser.vlm_doc_parser import VLMDocParser
            callback(0.1, "Start to parse Word document (VLM mode).")
            vlm_config = _resolve_vlm_config(parser_config)
            if not vlm_config.get("vlm_api_key") or not vlm_config.get("markdown_api_key"):
                callback(-1, "VLM config missing api_key. Set parser_config.vlm_config or backend env vars.")
                return []
            vlm_parser = VLMDocParser(api_config=vlm_config)
            file_bytes = binary if binary else open(filename, "rb").read()
            callback(0.2, "Rendering pages and calling VLM for layout analysis...")
            markdown_str, figures_dict, captions_dict = vlm_parser(file_bytes, file_type="word")

            if not markdown_str:
                callback(-1, "VLM parsing returned empty result for Word document.")
                return []

            chunk_token_num = _vlm_chunk_token_num(parser_config)
            markdown_str = _replace_vlm_placeholders(
                markdown_str,
                figures_dict,
                captions_dict,
                getattr(vlm_parser, "figure_order", None),
            )
            _export_vlm_assets_for_image_recall(filename, figures_dict, captions_dict, binary=binary)
            sections, tables = _build_vlm_sections(markdown_str, chunk_token_num)
            res = tokenize_table(tables, doc, is_english)
            section_images = _build_vlm_section_images(sections, figures_dict)
            chunks, images = _merge_vlm_sections_with_images(sections, section_images, chunk_token_num)
            res.extend(tokenize_chunks_with_images(chunks, doc, is_english, images))
            callback(0.8, "Finish VLM parsing (Word).")
            return res

        callback(0.1, "Start to parse.")

        try:
            vision_model = LLMBundle(kwargs["tenant_id"], LLMType.IMAGE2TEXT)
            callback(0.15, "Visual model detected. Attempting to enhance figure extraction...")
        except Exception:
            vision_model = None

        sections, tables = Docx()(filename, binary)

        if vision_model:
            figures_data = vision_figure_parser_figure_data_wraper(sections)
            try:
                docx_vision_parser = VisionFigureParser(vision_model=vision_model, figures_data=figures_data, **kwargs)
                boosted_figures = docx_vision_parser(callback=callback)
                tables.extend(boosted_figures)
            except Exception as e:
                callback(0.6, f"Visual model error: {e}. Skipping figure parsing enhancement.")

        res = tokenize_table(tables, doc, is_english)
        callback(0.8, "Finish parsing.")

        st = timer()

        chunks, images = naive_merge_docx(
            sections, int(parser_config.get(
                "chunk_token_num", 128)), parser_config.get(
                "delimiter", "\n!?。；！？"))

        if kwargs.get("section_only", False):
            return chunks

        res.extend(tokenize_chunks_with_images(chunks, doc, is_english, images))
        logging.info("naive_merge({}): {}".format(filename, timer() - st))
        return res

    elif re.search(r"\.pdf$", filename, re.IGNORECASE):
        layout_recognizer = parser_config.get("layout_recognize", "DeepDOC")
        if isinstance(layout_recognizer, bool):
            layout_recognizer = "DeepDOC" if layout_recognizer else "Plain Text"
        callback(0.1, "Start to parse.")

        if layout_recognizer == "DeepDOC":
            pdf_parser = Pdf()

            try:
                vision_model = LLMBundle(kwargs["tenant_id"], LLMType.IMAGE2TEXT)
                callback(0.15, "Visual model detected. Attempting to enhance figure extraction...")
            except Exception:
                vision_model = None

            if vision_model:
                sections, tables, figures = pdf_parser(filename if not binary else binary, from_page=from_page, to_page=to_page, callback=callback, separate_tables_figures=True)
                callback(0.5, "Basic parsing complete. Proceeding with figure enhancement...")
                try:
                    pdf_vision_parser = VisionFigureParser(vision_model=vision_model, figures_data=figures, **kwargs)
                    boosted_figures = pdf_vision_parser(callback=callback)
                    tables.extend(boosted_figures)
                except Exception as e:
                    callback(0.6, f"Visual model error: {e}. Skipping figure parsing enhancement.")
                    tables.extend(figures)
            else:
                sections, tables = pdf_parser(filename if not binary else binary, from_page=from_page, to_page=to_page, callback=callback)

            res = tokenize_table(tables, doc, is_english)
            callback(0.8, "Finish parsing.")

        elif layout_recognizer == "VLM":
            from deepdoc.parser.vlm_doc_parser import VLMDocParser
            callback(0.1, "Start to parse (VLM mode).")
            vlm_config = _resolve_vlm_config(parser_config)
            if not vlm_config.get("vlm_api_key") or not vlm_config.get("markdown_api_key"):
                callback(-1, "VLM config missing api_key. Set parser_config.vlm_config or backend env vars.")
                return []
            vlm_parser = VLMDocParser(api_config=vlm_config)
            file_bytes = binary if binary else open(filename, "rb").read()
            callback(0.2, "Rendering pages and calling VLM for layout analysis...")
            markdown_str, figures_dict, captions_dict = vlm_parser(
                file_bytes, file_type="pdf", from_page=from_page, to_page=to_page)

            if not markdown_str:
                callback(-1, "VLM parsing returned empty result.")
                return []

            chunk_token_num = _vlm_chunk_token_num(parser_config)
            markdown_str = _replace_vlm_placeholders(
                markdown_str,
                figures_dict,
                captions_dict,
                getattr(vlm_parser, "figure_order", None),
            )
            _export_vlm_assets_for_image_recall(filename, figures_dict, captions_dict, binary=binary)
            sections, tables = _build_vlm_sections(markdown_str, chunk_token_num)
            res = tokenize_table(tables, doc, is_english)
            section_images = _build_vlm_section_images(sections, figures_dict)
            chunks, images = _merge_vlm_sections_with_images(sections, section_images, chunk_token_num)
            res.extend(tokenize_chunks_with_images(chunks, doc, is_english, images))
            callback(0.8, "Finish VLM parsing.")
            return res

        else:
            if layout_recognizer == "Plain Text":
                pdf_parser = PlainParser()
            else:
                vision_model = LLMBundle(kwargs["tenant_id"], LLMType.IMAGE2TEXT, llm_name=layout_recognizer, lang=lang)
                pdf_parser = VisionParser(vision_model=vision_model, **kwargs)

            sections, tables = pdf_parser(filename if not binary else binary, from_page=from_page, to_page=to_page,
                                          callback=callback)
            res = tokenize_table(tables, doc, is_english)
            callback(0.8, "Finish parsing.")

    elif re.search(r"\.(csv|xlsx?)$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        excel_parser = ExcelParser()
        if parser_config.get("html4excel"):
            sections = [(_, "") for _ in excel_parser.html(binary, 12) if _]
        else:
            sections = [(_, "") for _ in excel_parser(binary) if _]

    elif re.search(r"\.(txt|py|js|java|c|cpp|h|php|go|ts|sh|cs|kt|sql)$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        sections = TxtParser()(filename, binary,
                               parser_config.get("chunk_token_num", 128),
                               parser_config.get("delimiter", "\n!?;。；！？"))
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.(md|markdown)$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        markdown_parser = Markdown(int(parser_config.get("chunk_token_num", 128)))
        sections, tables = markdown_parser(filename, binary)
        
        # Process images for each section
        section_images = []
        for section_text, _ in sections:
            images = markdown_parser.get_pictures(section_text) if section_text else None
            if images:
                # If multiple images found, combine them using concat_img
                combined_image = reduce(concat_img, images) if len(images) > 1 else images[0]
                section_images.append(combined_image)
            else:
                section_images.append(None)
                
        res = tokenize_table(tables, doc, is_english)
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.(htm|html)$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        sections = HtmlParser()(filename, binary)
        sections = [(_, "") for _ in sections if _]
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.json$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        chunk_token_num = int(parser_config.get("chunk_token_num", 128))
        sections = JsonParser(chunk_token_num)(binary)
        sections = [(_, "") for _ in sections if _]
        callback(0.8, "Finish parsing.")

    elif re.search(r"\.doc$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        binary = BytesIO(binary)
        doc_parsed = parser.from_buffer(binary)
        if doc_parsed.get('content', None) is not None:
            sections = doc_parsed['content'].split('\n')
            sections = [(_, "") for _ in sections if _]
            callback(0.8, "Finish parsing.")
        else:
            callback(0.8, f"tika.parser got empty content from {filename}.")
            logging.warning(f"tika.parser got empty content from {filename}.")
            return []

    else:
        raise NotImplementedError(
            "file type not supported yet(pdf, xlsx, doc, docx, txt supported)")

    st = timer()
    if section_images:
        # if all images are None, set section_images to None
        if all(image is None for image in section_images):
            section_images = None

    if section_images:
        chunks, images = naive_merge_with_images(sections, section_images,
                                        int(parser_config.get(
                                            "chunk_token_num", 128)), parser_config.get(
                                            "delimiter", "\n!?。；！？"))
        if kwargs.get("section_only", False):
            return chunks
        
        res.extend(tokenize_chunks_with_images(chunks, doc, is_english, images))
    else:
        chunks = naive_merge(
            sections, int(parser_config.get(
                "chunk_token_num", 128)), parser_config.get(
                "delimiter", "\n!?。；！？"))
        if kwargs.get("section_only", False):
            return chunks

        res.extend(tokenize_chunks(chunks, doc, is_english, pdf_parser))
    
    logging.info("naive_merge({}): {}".format(filename, timer() - st))
    return res


if __name__ == "__main__":
    import sys

    def dummy(prog=None, msg=""):
        pass

    chunk(sys.argv[1], from_page=0, to_page=10, callback=dummy)
